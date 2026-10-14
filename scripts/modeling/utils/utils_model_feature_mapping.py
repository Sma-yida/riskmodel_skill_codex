"""模型评估阶段的风控变量名称映射。

读取最终入模特征，匹配“数据变量名 -> 风控变量名 -> 中文名”，并输出给部署测试使用。
所有可修改路径和表头均集中在本文件顶部；不依赖 model_eval_config.json。
"""

import os
from pathlib import Path

import pandas as pd


# ==================== 用户可修改配置 ====================
# 映射表目录：优先用环境变量 RISKMODEL_MAPPING_DIR（skill 按本机设置传入），未设置时用工作目录下的 data/mapping/
MAPPING_DIR = Path(os.environ.get("RISKMODEL_MAPPING_DIR") or "data/mapping")
MAPPING_FILES = [
    {"path": str(MAPPING_DIR / "三方总表20260624.csv"), "name": "三方映射表"},
    {"path": str(MAPPING_DIR / "人行总表20260119.csv"), "name": "人行映射表"},
]
MODEL_FEATURE_TABLE_PATH = str(MAPPING_DIR / "模型特征总表_1783477836.csv")

DATA_VARIABLE_COLUMN = "数仓变量名"
RISK_VARIABLE_COLUMN = "风控变量名"
CHINESE_NAME_COLUMN = "变量中文名"
MODEL_CHINESE_NAME_COLUMN = "模型风控中文名"
MAPPING_RESULT_FILE = "haina_mapping_res.csv"
FK_TEST_FILE = "test_pred_100_FK.csv"
API_PARAMETER_FILE = "API参数配置模板.xlsx"
UNMAPPED_SOURCE = "映射关系未找到，待补充"
MODEL_TABLE_SOURCE = "模型补充表"
UNMAPPED_NOTE = "未找到映射关系"

# API 参数模板的固定列和字段取值。部署口径有变化时，只需在此处调整。
API_PARAMETER_COLUMNS = [
    "参数名称", "别名", "参数中文名", "参数使用说明", "父节点",
    "是否必填", "是否单条", "参数值类型", "参数类型", "备注",
]
REQUEST_REQUIRED = 1
REQUEST_VALUE_TYPE = "string"
REQUEST_PARAMETER_TYPE = "request_schema"
RESPONSE_PARAMETER_NAME = "body.data.probability(1)"
RESPONSE_VALUE_TYPE = "double"
RESPONSE_PARAMETER_TYPE = "response_schema"


def _clean_value(value):
    """统一读取到的文本；空值返回空字符串。"""
    if pd.isna(value):
        return ""
    return str(value).strip()


def _read_mapping_file(path, source_name):
    """读取三方/人行映射表；缺路径或缺表头只打印提示并返回空表。"""
    path = Path(path)
    if not path.exists():
        print(f"风控变量映射提示：{source_name}不存在，已跳过：{path}")
        return pd.DataFrame()

    try:
        frame = pd.read_csv(path, dtype=str)
    except Exception as error:
        print(f"风控变量映射提示：读取{source_name}失败，已跳过：{error}")
        return pd.DataFrame()

    required = {DATA_VARIABLE_COLUMN, RISK_VARIABLE_COLUMN, CHINESE_NAME_COLUMN}
    missing = required - set(frame.columns)
    if missing:
        print(f"风控变量映射提示：{source_name}缺少表头{sorted(missing)}，已跳过")
        return pd.DataFrame()
    return frame


def _build_source_mapping():
    """按 MAPPING_FILES 顺序建立映射；第一张表优先，冲突仅提示不覆盖。"""
    mapping = {}
    for item in MAPPING_FILES:
        source_name = item["name"]
        frame = _read_mapping_file(item["path"], source_name)
        for _, row in frame.iterrows():
            data_name = _clean_value(row[DATA_VARIABLE_COLUMN])
            risk_name = _clean_value(row[RISK_VARIABLE_COLUMN])
            chinese_name = _clean_value(row[CHINESE_NAME_COLUMN])
            if not data_name:
                continue
            new_value = (risk_name, chinese_name, source_name)
            if data_name in mapping:
                if mapping[data_name][:2] != new_value[:2]:
                    print(
                        f"风控变量映射冲突：{data_name} 已采用{mapping[data_name][2]}"
                        f"={mapping[data_name][:2]}，忽略{source_name}={new_value[:2]}"
                    )
                continue
            mapping[data_name] = new_value
    return mapping


def _build_model_table_mapping():
    """读取模型补充表：仅支持“风控变量名 -> 中文名”的同名兜底。"""
    path = Path(MODEL_FEATURE_TABLE_PATH)
    if not path.exists():
        print(f"风控变量映射提示：模型补充表不存在，已跳过：{path}")
        return {}

    try:
        frame = pd.read_csv(path, dtype=str)
    except Exception as error:
        print(f"风控变量映射提示：读取模型补充表失败，已跳过：{error}")
        return {}

    required = {RISK_VARIABLE_COLUMN, MODEL_CHINESE_NAME_COLUMN }
    missing = required - set(frame.columns)
    if missing:
        print(f"风控变量映射提示：模型补充表缺少表头{sorted(missing)}，已跳过")
        return {}

    result = {}
    for _, row in frame.iterrows():
        risk_name = _clean_value(row[RISK_VARIABLE_COLUMN])
        chinese_name = _clean_value(row[MODEL_CHINESE_NAME_COLUMN ])
        if risk_name and risk_name not in result:
            result[risk_name] = chinese_name
    return result


def build_feature_mapping(model_features):
    """输出最终入模特征的风控名称、中文名和映射来源。"""
    source_mapping = _build_source_mapping()
    model_table_mapping = _build_model_table_mapping()
    rows = []
    unmatched_features = []

    for feature in model_features:
        feature = str(feature)
        if feature in source_mapping:
            risk_name, chinese_name, source_name = source_mapping[feature]
        elif feature in model_table_mapping:
            risk_name, chinese_name, source_name = feature, model_table_mapping[feature], MODEL_TABLE_SOURCE
        else:
            risk_name, chinese_name, source_name = "", "", UNMAPPED_SOURCE
            unmatched_features.append(feature)
        rows.append({
            "特征名称": feature,
            "风控变量名称": risk_name,
            "特征中文名称": chinese_name,
            "映射来源": source_name,
        })
    if unmatched_features:
        print(
            "风控变量映射提示：以下入模特征未找到映射"
            f"（共{len(unmatched_features)}个）：{', '.join(unmatched_features)}"
        )
    return pd.DataFrame(rows)


def _rename_test_columns(test_frame, mapping_result, model_features):
    """仅重命名已匹配的入模特征；重复目标列名时保留原列名。"""
    mapping_by_feature = mapping_result.set_index("特征名称")
    selected_features = [feature for feature in model_features if feature in test_frame.columns]
    missing_features = [feature for feature in model_features if feature not in test_frame.columns]
    if missing_features:
        print(f"风控变量映射提示：测试样本缺少入模特征，已跳过：{missing_features}")

    rename_map = {}
    used_names = set()
    for feature in selected_features:
        risk_name = _clean_value(mapping_by_feature.loc[feature, "风控变量名称"])
        target_name = risk_name or feature
        if target_name in used_names:
            print(f"风控变量映射提示：风控变量名重复 {target_name}，保留原列名：{feature}")
            target_name = feature
        rename_map[feature] = target_name
        used_names.add(target_name)
    return test_frame[selected_features].rename(columns=rename_map)


def _write_api_parameter_template(mapping_result, retrain_data_dir):
    """按部署模板生成 API 参数配置 Excel，不改变映射和测试样本本身。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as error:
        print(f"风控变量映射提示：缺少 Excel 依赖，未生成API参数模板：{error}")
        return None

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(API_PARAMETER_COLUMNS)

    # 每个入模变量都是 API 请求字段；映射不到时仅在备注中标识，字段仍保留。
    for _, row in mapping_result.iterrows():
        risk_name = _clean_value(row["风控变量名称"])
        source_name = _clean_value(row["映射来源"])
        worksheet.append([
            _clean_value(row["特征名称"]),
            risk_name,
            _clean_value(row["特征中文名称"]),
            None,
            None,
            REQUEST_REQUIRED,
            None,
            REQUEST_VALUE_TYPE,
            REQUEST_PARAMETER_TYPE,
            UNMAPPED_NOTE if source_name == UNMAPPED_SOURCE or not risk_name else None,
        ])

    # 最后一行固定为模型概率输出字段。
    worksheet.append([
        RESPONSE_PARAMETER_NAME,
        None,
        None,
        None,
        None,
        REQUEST_REQUIRED,
        None,
        RESPONSE_VALUE_TYPE,
        RESPONSE_PARAMETER_TYPE,
        None,
    ])

    # 无原始 Excel 模板文件时，按截图建立稳定、可直接交付的表头和输出行样式。
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    response_fill = PatternFill("solid", fgColor="E2F0D9")
    thin_gray = Side(style="thin", color="B7B7B7")
    table_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for cell in worksheet[1]:
        cell.font = Font(name="宋体", bold=True, size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = table_border

    response_row = worksheet.max_row
    for row in worksheet.iter_rows(min_row=2, max_row=response_row, max_col=len(API_PARAMETER_COLUMNS)):
        for cell in row:
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="center")
            cell.border = table_border
    for column_index in (6, 7, 8, 9):
        for cell in worksheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2, max_row=response_row):
            cell[0].alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[response_row]:
        cell.fill = response_fill

    column_widths = [34, 30, 36, 36, 16, 12, 12, 16, 20, 24]
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    worksheet.row_dimensions[1].height = 24
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:J{response_row}"

    output_path = Path(retrain_data_dir) / API_PARAMETER_FILE
    workbook.save(output_path)
    print(f"API参数配置模板已保存：{output_path}")
    return output_path


def generate_risk_variable_mapping(model_features, retrain_data_dir, test_pred_path=None, prediction_column="prob_"):
    """生成映射结果和仅含风控字段的100条测试样本。

    不重新加载模型或重新预测；test_pred_100_FK.csv 仅改变列名并删除预测概率列。
    """
    retrain_data_dir = Path(retrain_data_dir)
    retrain_data_dir.mkdir(parents=True, exist_ok=True)
    model_features = [str(feature) for feature in model_features]

    mapping_result = build_feature_mapping(model_features)
    mapping_path = retrain_data_dir / MAPPING_RESULT_FILE
    mapping_result.to_csv(mapping_path, index=False, encoding="utf-8-sig")
    print(f"风控变量映射结果已保存：{mapping_path}")
    try:
        _write_api_parameter_template(mapping_result, retrain_data_dir)
    except Exception as error:
        print(f"风控变量映射提示：生成API参数模板失败，已跳过：{error}")

    test_pred_path = Path(test_pred_path or retrain_data_dir / "test_pred_100.csv")
    fk_test_path = retrain_data_dir / FK_TEST_FILE
    if not test_pred_path.exists():
        print(f"风控变量映射提示：测试样本不存在，未生成FK文件：{test_pred_path}")
        return mapping_result, None

    try:
        test_frame = pd.read_csv(test_pred_path)
        if prediction_column in test_frame.columns:
            test_frame = test_frame.drop(columns=[prediction_column])
        fk_test_frame = _rename_test_columns(test_frame, mapping_result, model_features)
        fk_test_frame.to_csv(fk_test_path, index=False, encoding="utf-8-sig")
        print(f"风控字段测试样本已保存：{fk_test_path}")
        return mapping_result, fk_test_frame
    except Exception as error:
        print(f"风控变量映射提示：生成FK测试样本失败，已跳过：{error}")
        return mapping_result, None
